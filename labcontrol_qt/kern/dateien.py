"""
LabControl - Messdateien einlesen
=================================

Liest die Laufdatei eines Geraets in ein rechteckiges Raster: erste Zeile die
Spaltenueberschriften, darunter je Probe eine Zeile. Wie lims_db.py ist diese
Datei frei von jeder GUI-Abhaengigkeit und dadurch automatisiert testbar.

Die Geraete liefern sehr unterschiedliche Dateien - Aufgabe dieses Moduls ist
nur, sie ueberhaupt aufzuspannen. Interpretiert wird hier nichts: Werte
bleiben als Text stehen, so wie sie in der Datei stehen. Das Umrechnen in ein
einheitliches Format und die Pruefungen kommen spaeter und arbeiten auf
diesem Raster.

Warum das Trennzeichen aus dem Inhalt kommt und nicht aus der Endung
---------------------------------------------------------------------
Die ICP-MS-Laufdatei heisst .csv, ist aber Tab-getrennt. Umgekehrt gibt es
.txt-Dateien mit Semikolon. Die Endung sagt also nichts. Erkannt wird am
Kopfzeileninhalt - und zwar an der Kopfzeile, weil in den Datenzeilen das
Dezimalkomma sonst als Trennzeichen durchginge.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import os

# Reihenfolge = Vorrang bei Gleichstand.
TRENNZEICHEN = ("\t", ";", "|", ",")

# utf-8-sig zuerst: die Geraetedateien tragen oft ein BOM. cp1252 faengt die
# Windows-Dateien mit Umlauten ab, latin-1 als letzte Rettung - das dekodiert
# jedes Byte und laesst den Einlesevorgang nie an der Kodierung scheitern.
KODIERUNGEN = ("utf-8-sig", "utf-8", "cp1252", "latin-1")

TABELLENBLATT_ENDUNGEN = (".xlsx", ".xlsm", ".xltx")

# Schutz gegen eine versehentlich riesige Datei: die Anzeige wuerde sonst
# minutenlang haengen. Wird die Grenze erreicht, sagt die Laufdatei das.
MAX_ZEILEN = 50000


class Laufdatei:
    """Das aufgespannte Raster einer Messdatei."""

    def __init__(self, pfad, spalten, zeilen, kodierung="", trennzeichen="",
                 abgeschnitten=False):
        self.pfad = pfad
        self.spalten = spalten
        self.zeilen = zeilen
        self.kodierung = kodierung
        self.trennzeichen = trennzeichen
        self.abgeschnitten = abgeschnitten

    @property
    def name(self) -> str:
        return os.path.basename(self.pfad)

    @property
    def zeilenzahl(self) -> int:
        return len(self.zeilen)

    @property
    def spaltenzahl(self) -> int:
        return len(self.spalten)

    def herkunft(self) -> str:
        """Kurzbeschreibung, wie die Datei gelesen wurde."""
        if self.trennzeichen:
            benennung = {"\t": "Tabulator", ";": "Semikolon", ",": "Komma",
                         "|": "Pipe"}.get(self.trennzeichen, self.trennzeichen)
            return f"Text, {benennung}-getrennt, {self.kodierung}"
        return "Tabellenblatt (xlsx)"


def erkenne_kodierung(rohdaten: bytes) -> str:
    """Erste Kodierung, mit der sich die Daten vollstaendig lesen lassen."""
    for kodierung in KODIERUNGEN:
        try:
            rohdaten.decode(kodierung)
            return kodierung
        except UnicodeDecodeError:
            continue
    return "latin-1"                      # dekodiert jedes Byte


def erkenne_trennzeichen(kopfzeile: str) -> str:
    """Bestimmt das Trennzeichen anhand der Kopfzeile.

    Bewusst nur die Kopfzeile: in den Datenzeilen kaeme sonst das
    Dezimalkomma als Trennzeichen in Betracht und wuerde jede Zahl zerlegen.
    """
    treffer = [(kopfzeile.count(zeichen), -platz, zeichen)
               for platz, zeichen in enumerate(TRENNZEICHEN)]
    anzahl, _, zeichen = max(treffer)
    return zeichen if anzahl else "\t"    # eine Spalte: Wahl egal


def _als_text(wert) -> str:
    """Stellt einen Zellwert aus einem Tabellenblatt als Text dar."""
    if wert is None:
        return ""
    if isinstance(wert, dt.datetime):
        return wert.strftime("%d.%m.%Y %H:%M:%S")
    if isinstance(wert, dt.date):
        return wert.strftime("%d.%m.%Y")
    if isinstance(wert, float) and wert.is_integer():
        return str(int(wert))             # 4.0 aus Excel ist eine 4
    return str(wert)


def _rechteckig(kopf: list[str], zeilen: list[list[str]]) -> tuple[list, list]:
    """Bringt Kopf und Zeilen auf dieselbe Breite.

    Ragged Dateien gibt es haeufiger als man denkt - eine zu kurze Zeile darf
    die Anzeige nicht verschieben, eine zu lange keine Werte verschlucken.
    """
    breite = max([len(kopf)] + [len(zeile) for zeile in zeilen] or [0])
    kopf = list(kopf) + [f"Spalte {n + 1}" for n in range(len(kopf), breite)]
    kopf = [name.strip() or f"Spalte {n + 1}" for n, name in enumerate(kopf)]
    zeilen = [list(zeile) + [""] * (breite - len(zeile)) for zeile in zeilen]
    return kopf, zeilen


def _lies_text(pfad: str) -> Laufdatei:
    with open(pfad, "rb") as datei:
        rohdaten = datei.read()
    if not rohdaten.strip():
        raise ValueError(f"Die Datei '{os.path.basename(pfad)}' ist leer.")

    kodierung = erkenne_kodierung(rohdaten)
    text = rohdaten.decode(kodierung)
    erste_zeile = next((z for z in text.splitlines() if z.strip()), "")
    trennzeichen = erkenne_trennzeichen(erste_zeile)

    leser = csv.reader(io.StringIO(text, newline=""), delimiter=trennzeichen)
    alle = [zeile for zeile in leser if any(feld.strip() for feld in zeile)]
    if not alle:
        raise ValueError(f"Die Datei '{os.path.basename(pfad)}' enthaelt keine Zeilen.")

    kopf, zeilen = alle[0], alle[1:]
    abgeschnitten = len(zeilen) > MAX_ZEILEN
    kopf, zeilen = _rechteckig(kopf, zeilen[:MAX_ZEILEN])
    return Laufdatei(pfad, kopf, zeilen, kodierung, trennzeichen, abgeschnitten)


def _lies_tabellenblatt(pfad: str) -> Laufdatei:
    try:
        import openpyxl
    except ImportError:
        raise ValueError(
            "Zum Lesen von xlsx-Dateien fehlt openpyxl im Programm."
        )
    try:
        mappe = openpyxl.load_workbook(pfad, read_only=True, data_only=True)
    except Exception as fehler:
        raise ValueError(f"'{os.path.basename(pfad)}' ist nicht lesbar: {fehler}")

    try:
        blatt = mappe[mappe.sheetnames[0]]
        alle = []
        for zeile in blatt.iter_rows(values_only=True):
            felder = [_als_text(zelle) for zelle in zeile]
            if any(feld.strip() for feld in felder):
                alle.append(felder)
            if len(alle) > MAX_ZEILEN + 1:
                break
    finally:
        mappe.close()

    if not alle:
        raise ValueError(f"Das Blatt in '{os.path.basename(pfad)}' ist leer.")

    kopf, zeilen = alle[0], alle[1:]
    abgeschnitten = len(zeilen) > MAX_ZEILEN
    kopf, zeilen = _rechteckig(kopf, zeilen[:MAX_ZEILEN])
    return Laufdatei(pfad, kopf, zeilen, abgeschnitten=abgeschnitten)


def lies(pfad: str) -> Laufdatei:
    """Liest eine Messdatei und gibt sie als Raster zurueck."""
    if not os.path.isfile(pfad):
        raise ValueError(f"Die Datei '{pfad}' gibt es nicht.")
    if os.path.splitext(pfad)[1].lower() in TABELLENBLATT_ENDUNGEN:
        return _lies_tabellenblatt(pfad)
    return _lies_text(pfad)
